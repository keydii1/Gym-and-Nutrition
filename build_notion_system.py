import csv
import openpyxl
import os

wb = openpyxl.load_workbook('Personal_Bulking_Tracker_2026.xlsx', data_only=True)
out_dir = 'notion_templates'
os.makedirs(out_dir, exist_ok=True)

# 1. EXERCISE DATABASE (Gym Hub)
ws07 = wb['07_WORKOUT_EXERCISE_DB']
exercises = []
for r in range(4, ws07.max_row + 1):
    ex_id = ws07.cell(r, 1).value
    name = ws07.cell(r, 2).value
    primary = ws07.cell(r, 3).value
    secondary = ws07.cell(r, 4).value
    equip = ws07.cell(r, 5).value
    ex_type = ws07.cell(r, 6).value
    diff = ws07.cell(r, 7).value
    guide = ws07.cell(r, 8).value
    rep_range = ws07.cell(r, 9).value
    alt = ws07.cell(r, 10).value
    safety = ws07.cell(r, 11).value
    if name:
        exercises.append({
            'Exercise Name': name,
            'Exercise ID': ex_id,
            'Primary Muscle': primary,
            'Secondary Muscle': secondary,
            'Equipment': equip,
            'Type': ex_type,
            'Difficulty': diff,
            'Target Rep Range': rep_range,
            'Alternative Exercise': alt,
            'Form Guide & Cues': guide,
            'Safety & Joint Notes': safety
        })

with open(f'{out_dir}/01_Gym_Exercise_Library.csv', 'w', encoding='utf-8-sig', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=list(exercises[0].keys()))
    writer.writeheader()
    writer.writerows(exercises)
print(f"Exported {len(exercises)} exercises to 01_Gym_Exercise_Library.csv")

# 2. WORKOUT ROUTINES / PROGRAMS (Gym Hub)
ws08 = wb['08_WEEKLY_WORKOUT_PLAN']
routines = []
for r in range(4, ws08.max_row + 1):
    week = ws08.cell(r, 1).value
    day = ws08.cell(r, 2).value
    session = ws08.cell(r, 3).value
    focus = ws08.cell(r, 4).value
    exercise = ws08.cell(r, 5).value
    order = ws08.cell(r, 6).value
    sets = ws08.cell(r, 7).value
    reps = ws08.cell(r, 8).value
    rir = ws08.cell(r, 9).value
    rest = ws08.cell(r, 10).value
    notes = ws08.cell(r, 12).value
    alt = ws08.cell(r, 13).value
    if exercise:
        routines.append({
            'Routine Item': f"{session} - {order}. {exercise}",
            'Session': session,
            'Day of Week': day,
            'Focus Muscle': focus,
            'Exercise': exercise,
            'Order': order,
            'Planned Sets': sets,
            'Target Reps': reps,
            'Target RIR': rir,
            'Rest Seconds': rest,
            'Progression Cues': notes,
            'Alternative': alt
        })

with open(f'{out_dir}/02_Gym_Workout_Routines.csv', 'w', encoding='utf-8-sig', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=list(routines[0].keys()))
    writer.writeheader()
    writer.writerows(routines)
print(f"Exported {len(routines)} routine items to 02_Gym_Workout_Routines.csv")

# 3. WORKOUT LOG STARTER (Gym Hub)
ws09 = wb['09_WORKOUT_LOG']
logs = []
for r in range(4, 16):
    date_val = ws09.cell(r, 1).value
    session = ws09.cell(r, 2).value
    exercise = ws09.cell(r, 3).value
    s_num = ws09.cell(r, 4).value
    load = ws09.cell(r, 5).value
    reps = ws09.cell(r, 6).value
    rir = ws09.cell(r, 7).value
    rest = ws09.cell(r, 8).value
    status = ws09.cell(r, 9).value
    pain = ws09.cell(r, 10).value
    notes = ws09.cell(r, 12).value
    if exercise:
        date_str = str(date_val).split(' ')[0] if date_val else '2026-10-20'
        vol = (float(load) * float(reps)) if load and reps else 0
        logs.append({
            'Log Entry': f"{date_str} - {session} - {exercise} (Set {s_num})",
            'Date': date_str,
            'Session': session,
            'Exercise': exercise,
            'Set': s_num,
            'Weight (kg)': load,
            'Reps': reps,
            'RIR': rir,
            'Rest (s)': rest,
            'Volume Load (kg)': vol,
            'Status': status,
            'Joint Pain Flag': pain,
            'Notes': notes
        })

with open(f'{out_dir}/03_Gym_Workout_Logs_Template.csv', 'w', encoding='utf-8-sig', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=list(logs[0].keys()))
    writer.writeheader()
    writer.writerows(logs)
print(f"Exported {len(logs)} workout log entries to 03_Gym_Workout_Logs_Template.csv")

# 4. FOOD DATABASE (Nutrition Hub)
ws03 = wb['03_FOOD_DATABASE']
foods = []
for r in range(5, ws03.max_row + 1):
    f_id = ws03.cell(r, 1).value
    name = ws03.cell(r, 2).value
    group = ws03.cell(r, 3).value
    state = ws03.cell(r, 4).value
    unit = ws03.cell(r, 5).value
    serving = ws03.cell(r, 6).value
    kcal = ws03.cell(r, 7).value
    p = ws03.cell(r, 8).value
    c = ws03.cell(r, 9).value
    fat = ws03.cell(r, 10).value
    fiber = ws03.cell(r, 11).value
    price = ws03.cell(r, 12).value
    src = ws03.cell(r, 13).value
    notes = ws03.cell(r, 14).value
    if name:
        foods.append({
            'Food Name': name,
            'Food ID': f_id,
            'Category': group,
            'State': state,
            'Default Unit': unit,
            'Standard Serving': serving,
            'Calories (per 100g)': kcal or 0,
            'Protein (g/100g)': p or 0,
            'Carbs (g/100g)': c or 0,
            'Fat (g/100g)': fat or 0,
            'Fiber (g/100g)': fiber or '',
            'Est Price (VNĐ/100g)': price or 0,
            'Source': src,
            'Nutrition Notes': notes
        })

with open(f'{out_dir}/04_Nutrition_Food_Database.csv', 'w', encoding='utf-8-sig', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=list(foods[0].keys()))
    writer.writeheader()
    writer.writerows(foods)
print(f"Exported {len(foods)} foods to 04_Nutrition_Food_Database.csv")

# 5. RECIPES DATABASE (Nutrition Hub)
ws04 = wb['04_RECIPE_DATABASE']
recipes = []
for r in range(5, ws04.max_row + 1):
    r_id = ws04.cell(r, 1).value
    name = ws04.cell(r, 2).value
    ingr = ws04.cell(r, 3).value
    kcal = ws04.cell(r, 4).value
    p = ws04.cell(r, 5).value
    c = ws04.cell(r, 6).value
    fat = ws04.cell(r, 7).value
    servings = ws04.cell(r, 8).value
    kcal_per = ws04.cell(r, 9).value
    p_per = ws04.cell(r, 10).value
    c_per = ws04.cell(r, 11).value
    fat_per = ws04.cell(r, 12).value
    cost = ws04.cell(r, 13).value
    time = ws04.cell(r, 14).value
    instr = ws04.cell(r, 15).value
    if name:
        recipes.append({
            'Recipe Name': name,
            'Recipe ID': r_id,
            'Ingredients List': ingr,
            'Total Calories': kcal,
            'Total Protein (g)': p,
            'Total Carbs (g)': c,
            'Total Fat (g)': fat,
            'Servings': servings,
            'Calories per Serving': kcal_per,
            'Protein per Serving (g)': p_per,
            'Carbs per Serving (g)': c_per,
            'Fat per Serving (g)': fat_per,
            'Cost per Serving (VNĐ)': cost,
            'Prep Time': time,
            'Cooking & Storage Instructions': instr
        })

with open(f'{out_dir}/05_Nutrition_Recipes_Database.csv', 'w', encoding='utf-8-sig', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=list(recipes[0].keys()))
    writer.writeheader()
    writer.writerows(recipes)
print(f"Exported {len(recipes)} recipes to 05_Nutrition_Recipes_Database.csv")

# 6. DAILY MEAL LOG TEMPLATE (Nutrition Hub)
ws06 = wb['06_DAILY_NUTRITION_LOG']
meals = []
for r in range(5, 17):
    d = ws06.cell(r, 1).value
    meal = ws06.cell(r, 2).value
    dish = ws06.cell(r, 3).value
    qty = ws06.cell(r, 4).value
    unit = ws06.cell(r, 5).value
    kcal = ws06.cell(r, 6).value
    p = ws06.cell(r, 7).value
    c = ws06.cell(r, 8).value
    fat = ws06.cell(r, 9).value
    cost = ws06.cell(r, 10).value
    conf = ws06.cell(r, 11).value
    notes = ws06.cell(r, 12).value
    if dish:
        date_str = str(d).split(' ')[0] if d else '2026-10-19'
        meals.append({
            'Meal Entry': f"{date_str} - {meal}: {dish}",
            'Date': date_str,
            'Meal Type': meal,
            'Food / Recipe': dish,
            'Amount': qty,
            'Unit': unit,
            'Calories (kcal)': round(kcal, 1) if kcal else 0,
            'Protein (g)': round(p, 1) if p else 0,
            'Carbs (g)': round(c, 1) if c else 0,
            'Fat (g)': round(fat, 1) if fat else 0,
            'Estimated Cost (VNĐ)': round(cost) if cost else 0,
            'Confidence': conf,
            'Notes': notes
        })

with open(f'{out_dir}/06_Nutrition_Daily_Meal_Log_Template.csv', 'w', encoding='utf-8-sig', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=list(meals[0].keys()))
    writer.writeheader()
    writer.writerows(meals)
print(f"Exported {len(meals)} meals to 06_Nutrition_Daily_Meal_Log_Template.csv")

# 7. BODY WEIGHT & MEASUREMENTS (Progress Hub)
body_logs = [
    {
        'Entry': '2026-10-19 - Baseline 19/10/2026',
        'Date': '2026-10-19',
        'Morning Weight (kg)': 50.3,
        'Waist (cm)': 68.0,
        'Chest (cm)': 82.0,
        'Arm (cm)': 25.5,
        'Thigh (cm)': 47.0,
        'Hip (cm)': 82.0,
        'Sleep (hours)': 7.5,
        'Appetite (1-5)': 3,
        'Energy (1-5)': 3,
        'Notes': 'Baseline đo InBody 19/10/2026. BMI 16.4 thiếu cân, PBF 7.8% rất thấp.'
    }
]
with open(f'{out_dir}/07_Body_Weight_Tracker.csv', 'w', encoding='utf-8-sig', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=list(body_logs[0].keys()))
    writer.writeheader()
    writer.writerows(body_logs)
print("Exported Body Weight starter to 07_Body_Weight_Tracker.csv")

# 8. INBODY HISTORY (Progress Hub)
inbody_logs = [
    {
        'Test Name': 'Lần 1 (Baseline InBody)',
        'Date': '2026-10-19',
        'Weight (kg)': 50.3,
        'BMI': 16.4,
        'SMM Muscle (kg)': 25.7,
        'Body Fat Mass (kg)': 3.9,
        'PBF Fat (%)': 7.8,
        'TBW Water (L)': 34.1,
        'Protein (kg)': 9.2,
        'Mineral (kg)': 3.06,
        'FFM Fat Free (kg)': 46.4,
        'BMR (kcal)': 1371,
        'InBody Score': 63,
        'Target Weight (kg)': 67.4,
        'Conditions': 'Buổi sáng, nhịn ăn, trước tập, sau khi đi vệ sinh',
        'Notes': 'Chỉ số ban đầu: BMI 16.4, PBF 7.8%. Mục tiêu dài hạn tăng SMM lên >30kg.'
    }
]
with open(f'{out_dir}/08_InBody_History.csv', 'w', encoding='utf-8-sig', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=list(inbody_logs[0].keys()))
    writer.writeheader()
    writer.writerows(inbody_logs)
print("Exported InBody starter to 08_InBody_History.csv")

print("\nALL NOTION CSV TEMPLATES SUCCESSFULLY BUILT!")
